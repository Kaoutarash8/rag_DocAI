"""Extracteurs DOCX / PPTX / XLSX / TXT / MD (sans dépendance à numpy ni à torch).

Chaque fonction retourne (nombre d'unités, générateur de PageContent), comme ocr.iter_pdf.
"""

import logging
import math
import re
import zipfile
from collections.abc import Iterator
from pathlib import Path

from app.core.config import get_settings
from app.services import vision
from app.services.content import HeadingStack, PageContent, Segment, rows_to_markdown

log = logging.getLogger(__name__)
GROUP_CHARS = 6000          # taille d'une unité pour les formats sans pages
XLSX_ROWS_PER_UNIT = 40
HEADING_STYLE = re.compile(r"^(?:heading|titre)\s*(\d)$", re.I)
MD_HEADING = re.compile(r"^(#{1,6})\s+(.*\S)\s*$")


def _describe(blob: bytes, mime: str, section: str, hint: str = "") -> Segment | None:
    if not vision.enabled() or len(blob) < 4_000:  # < 4 Ko : icône ou puce
        return None
    desc = vision.describe(blob, mime, section)
    return Segment(f"[Figure{hint}] {desc}", "figure", section) if desc else None


# ------------------------------------------------------------------ DOCX

def _docx_elements(path: Path) -> tuple[list[tuple], bool]:
    """Liste ordonnée d'éléments : ("heading", niveau, texte, page), ("para", texte, page),
    ("table", markdown, page), ("image", blob, mime, page)."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(path)
    out: list[tuple] = []
    page_lr = page_br = 0
    raw: list[tuple[int, int]] = []

    for child in doc.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, doc)
            lr = len(child.xpath(".//w:lastRenderedPageBreak"))
            br = len(child.xpath('.//w:br[@w:type="page"]'))
            page_lr += lr
            page_br += br
            idx = len(out)
            text = p.text.strip()
            m = HEADING_STYLE.match(p.style.name) if p.style is not None else None
            if text and m:
                out.append(("heading", int(m.group(1)), text, 0))
            elif text and p.style is not None and p.style.name.lower() in ("title", "titre"):
                out.append(("heading", 1, text, 0))
            elif text:
                out.append(("para", text, 0))
            if vision.enabled():
                for rid in child.xpath(".//a:blip/@r:embed"):
                    part = doc.part.related_parts.get(rid)
                    if part is not None:
                        out.append(("image", part.blob, part.content_type, 0))
            raw.append((page_lr, page_br))
            out_len = len(out)
            # on associe le numéro de page courant aux éléments ajoutés par ce paragraphe
            for j in range(idx, out_len):
                out[j] = out[j][:-1] + (len(raw) - 1,)  # indice temporaire dans `raw`
        elif tag == "tbl":
            t = Table(child, doc)
            rows = []
            for r in t.rows:
                cells, last = [], None
                for c in r.cells:
                    if c._tc is not last:  # cellules fusionnées : ne pas dupliquer
                        cells.append(c.text)
                    last = c._tc
                rows.append(cells)
            md = rows_to_markdown(rows)
            if md:
                raw.append((page_lr, page_br))
                out.append(("table", md, len(raw) - 1))

    has_pages = page_lr > 0 or page_br > 0
    use_lr = page_lr > 0  # lastRenderedPageBreak, plus fiable quand il est présent
    final = []
    for el in out:
        lr, br = raw[el[-1]]
        page = 1 + (lr if use_lr else br) if has_pages else None
        final.append(el[:-1] + (page,))
    return final, has_pages


def _group_docx(elements: list[tuple], has_pages: bool) -> list[list[tuple]]:
    groups: list[list[tuple]] = []
    cur: list[tuple] = []
    chars = 0
    last_page = None
    for el in elements:
        page = el[-1]
        boundary = (page != last_page) if has_pages else (chars >= GROUP_CHARS)
        if cur and boundary:
            groups.append(cur)
            cur, chars = [], 0
        cur.append(el)
        chars += len(el[1]) if isinstance(el[1], str) else 0
        last_page = page
    if cur:
        groups.append(cur)
    return groups


def iter_docx(path: Path) -> tuple[int, Iterator[PageContent]]:
    elements, has_pages = _docx_elements(path)
    groups = _group_docx(elements, has_pages)

    def gen() -> Iterator[PageContent]:
        stack = HeadingStack()
        for g in groups:
            segs: list[Segment] = []
            buf: list[str] = []
            section = stack.path

            def flush():
                nonlocal buf
                if buf:
                    segs.append(Segment("\n".join(buf), "text", section))
                    buf = []

            for el in g:
                kind = el[0]
                if kind == "heading":
                    flush()
                    stack.push(el[1], el[2])
                    section = stack.path
                    buf = [el[2]]
                elif kind == "para":
                    buf.append(el[1])
                elif kind == "table":
                    flush()
                    segs.append(Segment(el[1], "table", stack.path))
                elif kind == "image":
                    flush()
                    seg = _describe(el[1], el[2], stack.path)
                    if seg:
                        segs.append(seg)
            flush()
            yield PageContent(g[0][-1], segs)

    return len(groups), gen()


# ------------------------------------------------------------------ PPTX

def _pptx_shapes(shapes):
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for sh in shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _pptx_shapes(sh.shapes)
        else:
            yield sh


def _chart_markdown(chart) -> str:
    plot = chart.plots[0]
    cats = [str(c) for c in plot.categories]
    series = list(plot.series)
    rows = [[""] + [s.name for s in series]]
    for i, c in enumerate(cats):
        rows.append([c] + [list(s.values)[i] if i < len(list(s.values)) else "" for s in series])
    title = chart.chart_title.text_frame.text if chart.has_title else ""
    md = rows_to_markdown(rows)
    return (f"Graphique « {title} »\n" if title else "Graphique\n") + md if md else ""


def iter_pptx(path: Path) -> tuple[int, Iterator[PageContent]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(path)
    slides = list(prs.slides)

    def gen() -> Iterator[PageContent]:
        for n, slide in enumerate(slides, 1):
            title = ""
            try:
                if slide.shapes.title is not None:
                    title = slide.shapes.title.text_frame.text.strip()
            except Exception:  # noqa: BLE001
                pass
            section = title or f"Diapositive {n}"
            shapes = sorted(_pptx_shapes(slide.shapes),
                            key=lambda s: (s.top or 0, s.left or 0))
            text_lines, extra = [], []
            for sh in shapes:
                try:
                    if sh.has_text_frame and sh.text_frame.text.strip():
                        t = sh.text_frame.text.strip()
                        if t != title:
                            text_lines.append(t)
                    elif getattr(sh, "has_table", False) and sh.has_table:
                        md = rows_to_markdown([[c.text for c in r.cells] for r in sh.table.rows])
                        if md:
                            extra.append(Segment(md, "table", section))
                    elif getattr(sh, "has_chart", False) and sh.has_chart:
                        md = _chart_markdown(sh.chart)
                        if md:
                            extra.append(Segment(md, "figure", section))
                    elif sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                        seg = _describe(sh.image.blob, sh.image.content_type, section)
                        if seg:
                            extra.append(seg)
                except Exception as e:  # noqa: BLE001
                    log.warning("Forme ignorée (diapo %d) : %s", n, e)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                notes = slide.notes_slide.notes_text_frame.text.strip()
                if notes:
                    text_lines.append("Notes : " + notes)
            segs = []
            body = "\n".join(([title] if title else []) + text_lines)
            if body.strip():
                segs.append(Segment(body, "text", section))
            yield PageContent(n, segs + extra)

    return len(slides), gen()


# ------------------------------------------------------------------ XLSX

def iter_xlsx(path: Path) -> tuple[int, Iterator[PageContent]]:
    from openpyxl import load_workbook

    s = get_settings()
    wb = load_workbook(path, read_only=True, data_only=True)
    sheets = wb.worksheets
    total = sum(max(1, math.ceil(min(ws.max_row or 1, s.max_sheet_rows) / XLSX_ROWS_PER_UNIT))
                for ws in sheets)

    def gen() -> Iterator[PageContent]:
        try:
            for ws in sheets:
                section = f"Feuille {ws.title}"
                header, batch, seen_rows = None, [], 0
                for row in ws.iter_rows(values_only=True):
                    if not any(c is not None and str(c).strip() for c in row):
                        continue
                    seen_rows += 1
                    if seen_rows > s.max_sheet_rows:
                        log.warning("%s tronquée à %d lignes", section, s.max_sheet_rows)
                        break
                    if header is None:
                        header = list(row)
                        continue
                    batch.append(list(row))
                    if len(batch) >= XLSX_ROWS_PER_UNIT:
                        md = rows_to_markdown([header] + batch)
                        yield PageContent(None, [Segment(md, "table", section)])
                        batch = []
                if batch or header is not None:
                    md = rows_to_markdown([header] + batch) if header else ""
                    if md:
                        yield PageContent(None, [Segment(md, "table", section)])
        finally:
            wb.close()

    return total, gen()


# ------------------------------------------------------------------ TXT / MD

def read_text(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def iter_text(path: Path, markdown: bool) -> tuple[int, Iterator[PageContent]]:
    lines = read_text(path).splitlines()
    units: list[list[tuple[int, str]]] = []  # (niveau de titre ou 0, texte)
    cur, chars, in_code = [], 0, False
    for ln in lines:
        if markdown and ln.strip().startswith("```"):
            in_code = not in_code
        m = MD_HEADING.match(ln) if markdown and not in_code else None
        cur.append((len(m.group(1)) if m else 0, m.group(2) if m else ln))
        chars += len(ln)
        if chars >= GROUP_CHARS and not ln.strip():  # coupe sur ligne vide
            units.append(cur)
            cur, chars = [], 0
    if cur:
        units.append(cur)

    def gen() -> Iterator[PageContent]:
        stack = HeadingStack()
        for u in units:
            segs, buf, section = [], [], stack.path
            for level, text in u:
                if level:
                    if buf:
                        segs.append(Segment("\n".join(buf).strip(), "text", section))
                    stack.push(level, text)
                    section, buf = stack.path, [text]
                else:
                    buf.append(text)
            if "".join(buf).strip():
                segs.append(Segment("\n".join(buf).strip(), "text", section))
            yield PageContent(None, [sg for sg in segs if sg.text])

    return len(units), gen()


# ------------------------------------------------------------------ contrôle des archives Office

OFFICE_MARKERS = {"docx": "word/", "pptx": "ppt/", "xlsx": "xl/"}
MAX_UNCOMPRESSED = 1_000_000_000  # 1 Go : protège des « zip bombs »


def check_office_zip(path: Path, kind: str) -> bool:
    try:
        with zipfile.ZipFile(path) as z:
            infos = z.infolist()
            if sum(i.file_size for i in infos) > MAX_UNCOMPRESSED:
                return False
            names = z.namelist()
            return "[Content_Types].xml" in names and any(
                n.startswith(OFFICE_MARKERS[kind]) for n in names)
    except zipfile.BadZipFile:
        return False
