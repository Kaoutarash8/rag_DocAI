"""Tests des extracteurs : logique pure (titres, tableaux) et formats Office/texte."""
from collections import Counter


from app.services import office
from app.services.content import HeadingStack, font_profile, rows_to_markdown
from app.services.files import sniff_file_type
from app.services.ocr import build_segments, is_heading


def L(text, size=10.0, y=0.0, bold=False):
    return {"text": text, "size": size, "bold": bold, "y0": y, "x0": 0}


# ---- titres
def test_font_profile_detects_heading_levels():
    body, levels = font_profile(Counter({10.0: 9000, 18.0: 120, 14.0: 300, 10.5: 50}))
    assert body == 10.0 and levels == [18.0, 14.0]


def test_headings_become_sections_across_pages():
    stack = HeadingStack()
    body, levels = 10.0, [18.0, 14.0]
    p1 = [L("1 Introduction", 18, 0), L("texte un", 10, 1), L("1.1 Contexte", 14, 2), L("texte deux", 10, 3)]
    p2 = [L("suite du contexte", 10, 0), L("2 Méthode", 18, 1), L("texte trois", 10, 2)]
    s1 = build_segments(p1, [], body, levels, stack)
    s2 = build_segments(p2, [], body, levels, stack)
    assert [s.section for s in s1] == ["1 Introduction", "1 Introduction > 1.1 Contexte"]
    assert s2[0].section == "1 Introduction > 1.1 Contexte"   # section héritée de la page précédente
    assert s2[1].section == "2 Méthode"


def test_not_a_heading():
    assert is_heading(L("12"), 10, [18.0]) == 0                 # numéro de page
    assert is_heading(L("x" * 200, 18), 10, [18.0]) == 0        # trop long
    assert is_heading(L("Résultats", 10, bold=True), 10, [18.0]) == 2   # gras court
    assert is_heading(L("Une phrase en gras.", 10, bold=True), 10, [18.0]) == 0


def test_table_inserted_at_its_position():
    stack = HeadingStack()
    lines = [L("avant", 10, 0), L("après", 10, 50)]
    segs = build_segments(lines, [{"markdown": "| a | b |\n|---|---|\n| 1 | 2 |", "y0": 20}],
                          10.0, [], stack)
    assert [s.kind for s in segs] == ["text", "table", "text"]


def test_rows_to_markdown_escapes_and_pads():
    md = rows_to_markdown([["a|b", "c"], ["1"], [None, None]])
    assert md.splitlines()[0] == "| a\\|b | c |"
    assert md.splitlines()[2] == "| 1 |  |"
    assert rows_to_markdown([[None]]) == ""


# ---- formats
def _drain(res):
    total, gen = res
    pages = list(gen)
    return total, pages


def test_docx_headings_tables_pages(tmp_path):
    from docx import Document
    d = Document()
    d.add_heading("Rapport", 1)
    d.add_paragraph("Introduction du rapport.")
    d.add_heading("Chiffres", 2)
    t = d.add_table(rows=3, cols=2)
    for i, (a, b) in enumerate([("Produit", "CA"), ("A", "10"), ("B", "20")]):
        t.cell(i, 0).text, t.cell(i, 1).text = a, b
    d.add_paragraph("Conclusion.")
    path = tmp_path / "r.docx"
    d.save(path)
    total, pages = _drain(office.iter_docx(path))
    segs = [s for p in pages for s in p.segments]
    assert total == len(pages) >= 1
    table = next(s for s in segs if s.kind == "table")
    assert "| Produit | CA |" in table.text and table.section == "Rapport > Chiffres"
    assert any(s.section == "Rapport" and "Introduction" in s.text for s in segs)


def test_pptx_slides_tables_charts_notes(tmp_path):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Ventes 2025"
    cd = CategoryChartData()
    cd.categories = ["T1", "T2"]
    cd.add_series("CA", (10, 20))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), cd)
    s.notes_slide.notes_text_frame.text = "Penser à citer la source"
    path = tmp_path / "d.pptx"
    prs.save(path)
    total, pages = _drain(office.iter_pptx(path))
    assert total == 1 and pages[0].page == 1
    kinds = {sg.kind: sg for sg in pages[0].segments}
    assert "| T1 | 10" in kinds["figure"].text and kinds["figure"].section == "Ventes 2025"
    assert "Penser à citer la source" in kinds["text"].text


def test_xlsx_batches_repeat_header(tmp_path):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventes"
    ws.append(["Produit", "CA"])
    for i in range(95):
        ws.append([f"p{i}", i])
    path = tmp_path / "v.xlsx"
    wb.save(path)
    total, pages = _drain(office.iter_xlsx(path))
    assert len(pages) == 3 and all(p.page is None for p in pages)
    assert all(p.segments[0].text.startswith("| Produit | CA |") for p in pages)
    assert pages[0].segments[0].section == "Feuille Ventes"


def test_markdown_sections_and_code_fence(tmp_path):
    f = tmp_path / "n.md"
    f.write_text("# Titre\nintro\n## Sous-titre\ntexte\n```\n# pas un titre\n```\n", encoding="utf-8")
    _, pages = _drain(office.iter_text(f, markdown=True))
    sections = [s.section for s in pages[0].segments]
    assert sections == ["Titre", "Titre > Sous-titre"]
    assert "# pas un titre" in pages[0].segments[-1].text


def test_txt_latin1_fallback(tmp_path):
    f = tmp_path / "a.txt"
    f.write_bytes("café très bon".encode("cp1252"))
    _, pages = _drain(office.iter_text(f, markdown=False))
    assert "café" in pages[0].text


# ---- sniffing
def test_sniff_office_and_text(tmp_path):
    from docx import Document
    d = tmp_path / "a.docx"
    Document().save(d)
    assert sniff_file_type(d.read_bytes()[:16], ".docx", d) == "docx"
    assert sniff_file_type(d.read_bytes()[:16], ".xlsx", d) is None     # mauvaise extension
    fake = tmp_path / "f.docx"
    fake.write_bytes(b"PK\x03\x04" + b"junk" * 10)
    assert sniff_file_type(b"PK\x03\x04junk", ".docx", fake) is None    # zip invalide
    t = tmp_path / "t.md"
    t.write_text("# ok")
    assert sniff_file_type(b"# ok", ".md", t) == "md"
    b = tmp_path / "b.txt"
    b.write_bytes(b"abc\x00def")
    assert sniff_file_type(b"abc\x00def", ".txt", b) is None
